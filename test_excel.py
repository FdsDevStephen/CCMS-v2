"""Read the Excel results file."""
import openpyxl
wb = openpyxl.load_workbook("results/legal_extraction_results.xlsx")
ws = wb.active
print(f"Sheet: {ws.title}, Rows: {ws.max_row}, Cols: {ws.max_column}")

# Print headers
headers = [c.value for c in next(ws.iter_rows(max_row=1))]
print(f"\nHeaders: {headers}")

# Print all rows
print()
for row in ws.iter_rows(min_row=2, values_only=True):
    vals = [str(v)[:60] if v else "" for v in row]
    print(" | ".join(vals))
